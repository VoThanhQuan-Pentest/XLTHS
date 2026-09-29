import sys
from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Palette
    NAVY = RGBColor(10, 25, 47)       # #0A192F
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(246, 248, 250)
    DARK_CARD = RGBColor(17, 34, 64)   # #112240
    CARD_BG = RGBColor(255, 255, 255)
    TEXT_MAIN = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139)
    CYAN = RGBColor(0, 180, 216)      # #00B4D8
    TEAL = RGBColor(13, 148, 136)
    ORANGE = RGBColor(247, 127, 0)
    GREEN = RGBColor(16, 185, 129)
    RED = RGBColor(239, 68, 68)
    BORDER_CLR = RGBColor(226, 232, 240)
    
    def set_slide_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, subtitle_text, dark=False):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.name = "Arial"
        p.font.color.rgb = CYAN if dark else NAVY
        
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(13)
        p2.font.name = "Arial"
        p2.font.color.rgb = RGBColor(148, 163, 184) if dark else TEXT_MUTED

    def add_card(slide, x, y, w, h, bg_color=WHITE, border_color=BORDER_CLR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: CHÀO MỪNG & GIỚI THIỆU
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, NAVY)
    
    # Glow badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.8), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = DARK_CARD
    badge.line.color.rgb = CYAN
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.text = "BÁO CÁO GIỮA KỲ — XỬ LÝ TÍN HIỆU SỐ"
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = CYAN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Title & Subtitle
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PHÂN ĐOẠN TÍN HIỆU THÀNH\nTIẾNG NÓI VÀ KHOẢNG LẶNG"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.text = "So sánh thực nghiệm ba thuật toán: Tìm kiếm nhị phân, Histogram và Thống kê Gaussian"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    
    # 3 Methods Cards
    methods = [
        ("01. TÌM KIẾM NHỊ PHÂN", "Binary Search cân bằng diện tích nhầm lẫn hai lớp, khóa ngưỡng cố định.", CYAN),
        ("02. HISTOGRAM", "Phân tích 2 đỉnh năng lượng, tính ngưỡng thích nghi theo từng file WAV.", TEAL),
        ("03. THỐNG KÊ GAUSSIAN", "Mô hình hóa mật độ xác suất chuẩn, giải nghiệm Bayes tối ưu sai số.", ORANGE)
    ]
    for i, (m_title, m_desc, clr) in enumerate(methods):
        card = add_card(s1, 0.8 + i * 4.0, 4.3, 3.7, 1.6, DARK_CARD, clr)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = clr
        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        
    # Footer student info
    f_box = s1.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.6))
    tf = f_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Sinh viên: Võ Thanh Quân  |  Mã nguồn: 100% Python/Numpy  |  Thời lượng thuyết minh: 9 Phút"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: BIỂU ĐỒ KHỐI THUẬT TOÁN 1 (BINARY)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, LIGHT_BG)
    add_header(s2, "THUẬT TOÁN 1: TÌM KIẾM NHỊ PHÂN (BINARY SEARCH)", "Sơ đồ khối và nguyên lý xác định ngưỡng phân đoạn năng lượng")
    
    # 6-step Flowchart
    steps1 = [
        ("BƯỚC 1", "Dữ liệu huấn luyện & LAB", "Nạp 4 file TinHieuHuanLuyen và nhãn chuẩn tương ứng"),
        ("BƯỚC 2", "Trích xuất STE chuẩn hóa", "Chia khung 30 ms, bước nhảy 10 ms, chuẩn hóa STE về [0, 1]"),
        ("BƯỚC 3", "Tách 2 tập nhãn Sp & Sil", "Gộp 'v' + 'uv' thành Speech, 'sil' là Silence"),
        ("BƯỚC 4", "Khảo sát miền chồng lấn", "Xác định khoảng tìm kiếm [min(Speech), max(Silence)]"),
        ("BƯỚC 5", "Duyệt nhị phân chia đôi", "Cân bằng diện tích lỗi: P(Sil > mid) ≈ P(Sp < mid)"),
        ("BƯỚC 6", "Khóa ngưỡng & Lọc 200 ms", "Khóa T = 0.00156; loại khoảng lặng ảo ngắn hơn 200 ms")
    ]
    for i, (tag, title, desc) in enumerate(steps1):
        x = 0.8 + (i % 3) * 4.0
        y = 1.7 + (i // 3) * 2.3
        card = add_card(s2, x, y, 3.7, 2.0, WHITE)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = NAVY
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    # Bottom note banner
    banner = add_card(s2, 0.8, 6.4, 11.7, 0.65, RGBColor(238, 242, 255), CYAN)
    tf = banner.text_frame
    p = tf.paragraphs[0]
    p.text = "Ý NGHĨA KỸ THUẬT: Ngưỡng được khóa từ tập huấn luyện và dùng chung cho toàn bộ tập kiểm thử. Không sử dụng nhãn test để tinh chỉnh."
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY

    # ==========================================
    # SLIDE 3: BIỂU ĐỒ KẾT QUẢ THUẬT TOÁN 1
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, LIGHT_BG)
    add_header(s3, "KẾT QUẢ THỰC NGHIỆM: TÌM KIẾM NHỊ PHÂN TRÊN PHONE_F2", "Đồ thị 4 dải trung gian và kết quả phân đoạn cuối cùng")
    
    img_path = Path("ket_qua/phone_F2_binary.png")
    if img_path.exists():
        s3.shapes.add_picture(str(img_path), Inches(0.8), Inches(1.6), width=Inches(8.5))
        
    side_card = add_card(s3, 9.6, 1.6, 2.9, 5.3, WHITE)
    tf = side_card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "CẤU TRÚC 4 ĐỒ THỊ"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    sections = [
        ("Dải 1: Waveform", "Sóng âm thanh WAV và biên phân đoạn dọc."),
        ("Dải 2: STE chuẩn hóa", "Đường năng lượng và ngưỡng ngang T = 0.0016."),
        ("Dải 3: Mức năng lượng", "logSTE và logMA đo bằng dB."),
        ("Dải 4: Tần số F0", "F0 tự tương quan (Hz) và F0mean chuẩn LAB."),
        ("Quy ước đường kẻ", "• Đỏ nét đứt: Biên chuẩn LAB\n• Xanh nét liền: Biên thuật toán")
    ]
    for stitle, sdesc in sections:
        p1 = tf.add_paragraph()
        p1.text = stitle
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = CYAN
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: NHẬN XÉT BIỂU ĐỒ THUẬT TOÁN 1
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, LIGHT_BG)
    add_header(s4, "NHẬN XÉT CHI TIẾT: THUẬT TOÁN TÌM KIẾM NHỊ PHÂN", "Đánh giá định lượng sai số biên, vị trí và nguyên nhân trên đồ thị")
    
    # 3 Stat Cards
    metrics1 = [
        ("MAE TRÊN PHONE_F2", "45.0 ms", "RMSE = 57.0 ms", ORANGE),
        ("MAE TOÀN TẬP TEST", "20.0 ms", "Gộp trên 4 tệp kiểm thử", GREEN),
        ("SỐ LƯỢNG BIÊN", "2 Đúng / 1 Thừa / 0 Thiếu", "1 lỗi biên thừa tại đuôi", NAVY)
    ]
    for i, (mtitle, mval, msub, clr) in enumerate(metrics1):
        c = add_card(s4, 0.8 + i * 4.0, 1.6, 3.7, 1.3, WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = mtitle
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        p2 = tf.add_paragraph()
        p2.text = mval
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = clr
        p3 = tf.add_paragraph()
        p3.text = msub
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED
        
    # 2 Detailed Analysis Boxes
    c_left = add_card(s4, 0.8, 3.2, 5.7, 3.8, WHITE)
    tf = c_left.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "VỊ TRÍ VÀ MỨC ĐỘ SAI LỆCH BIÊN"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    items_left = [
        ("• Bắt đầu Speech (chuẩn 1.02 s):", "Thuật toán dự đoán 1.01 s (lệch -10 ms) -> Bám cực sát biên chuẩn."),
        ("• Kết thúc Speech (chuẩn 4.04 s):", "Thuật toán dự đoán 4.12 s (lệch +80 ms) -> Sai lệch vừa phải do âm đuôi suy giảm chậm."),
        ("• Biên thừa tại vị trí 4.76 s:", "Nhiễu nền cuối file vượt ngưỡng T = 0.0016, tạo 1 đoạn Speech giả ngắn.")
    ]
    for k, v in items_left:
        p1 = tf.add_paragraph()
        p1.text = k
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = CYAN
        p2 = tf.add_paragraph()
        p2.text = v
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    c_right = add_card(s4, 6.8, 3.2, 5.7, 3.8, WHITE)
    tf = c_right.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "NGUYÊN NHÂN & ĐÁNH GIÁ ĐẶC TRƯNG"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    items_right = [
        ("• Nguyên nhân sai lệch:", "Ngưỡng cố định T học từ tập training chung có thể hơi thấp so với mức nhiễu nền của riêng bản ghi phone_F2."),
        ("• Đường tần số cơ bản F0:", "F0 chỉ xuất hiện ở các khung hữu thanh, trung vị đạt 142.9 Hz (rất gần F0mean LAB 145.0 Hz)."),
        ("• Khoảng lặng ảo < 200 ms:", "Đã được loại bỏ triệt để, không phát sinh chia cắt sai lệch giữa câu.")
    ]
    for k, v in items_right:
        p1 = tf.add_paragraph()
        p1.text = k
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = TEAL
        p2 = tf.add_paragraph()
        p2.text = v
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 5: BIỂU ĐỒ KHỐI THUẬT TOÁN 2 (HISTOGRAM)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, LIGHT_BG)
    add_header(s5, "THUẬT TOÁN 2: HISTOGRAM MỨC NĂNG LƯỢNG (ADAPTIVE)", "Sơ đồ khối và cơ chế tìm ngưỡng thích nghi từ phân bố năng lượng")
    
    steps2 = [
        ("BƯỚC 1", "Tệp âm thanh kiểm thử", "Nhận tín hiệu WAV kiểm thử trực tiếp (hoàn toàn không cần nhãn LAB)"),
        ("BƯỚC 2", "Trích xuất STE chuẩn hóa", "Chia khung phân tích, tính toán mảng năng lượng normalized STE"),
        ("BƯỚC 3", "Lập Histogram 64 Bins", "Thống kê mật độ phân bố năng lượng trong miền [0, 1]"),
        ("BƯỚC 4", "Làm trơn Moving Average", "Lọc trơn 1D bằng Numpy để triệt tiêu các dao động nhiễu cục bộ"),
        ("BƯỚC 5", "Tìm 2 đỉnh cục bộ", "Phát hiện đỉnh năng lượng Silence (thấp) và đỉnh Speech (cao)"),
        ("BƯỚC 6", "Tính ngưỡng có trọng số", "T = (5*Peak_Sil + Peak_Sp) / 6; lọc khoảng lặng ảo < 200 ms")
    ]
    for i, (tag, title, desc) in enumerate(steps2):
        x = 0.8 + (i % 3) * 4.0
        y = 1.7 + (i // 3) * 2.3
        card = add_card(s5, x, y, 3.7, 2.0, WHITE)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEAL
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = NAVY
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    banner2 = add_card(s5, 0.8, 6.4, 11.7, 0.65, RGBColor(240, 253, 250), TEAL)
    tf = banner2.text_frame
    p = tf.paragraphs[0]
    p.text = "ĐẶC TRƯNG THÍCH NGHI: Ngưỡng được tự động xác định lại cho từng file test riêng biệt, tự thích ứng theo độ lớn âm lượng bản ghi."
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY

    # ==========================================
    # SLIDE 6: BIỂU ĐỒ KẾT QUẢ THUẬT TOÁN 2
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, LIGHT_BG)
    add_header(s6, "KẾT QUẢ THỰC NGHIỆM: HISTOGRAM TRÊN PHONE_F2", "Đồ thị 4 dải thể hiện hiện tượng phân mảnh tiếng nói do ngưỡng cao")
    
    img_path = Path("ket_qua/phone_F2_histogram.png")
    if img_path.exists():
        s6.shapes.add_picture(str(img_path), Inches(0.8), Inches(1.6), width=Inches(8.5))
        
    side_card2 = add_card(s6, 9.6, 1.6, 2.9, 5.3, WHITE)
    tf = side_card2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "THÔNG SỐ THỰC NGHIỆM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    sections2 = [
        ("Ngưỡng tự thích nghi", "T = 0.0572 (cao gấp ~35 lần so với Binary Search)."),
        ("Hiện tượng phân mảnh", "STE rơi dưới ngưỡng ở giữa câu, tạo 4 biên thừa."),
        ("Biên xanh vs Biên đỏ", "Các đoạn tiếng nói yếu bị ngắt thành Silence giả."),
        ("Dải tần số F0", "Vẫn đo được ổn định ở các phân đoạn hữu thanh."),
        ("Đánh giá tổng thể", "Thích nghi tốt với âm lượng lớn nhưng dễ xén tiếng nói nhỏ.")
    ]
    for stitle, sdesc in sections2:
        p1 = tf.add_paragraph()
        p1.text = stitle
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = TEAL
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 7: NHẬN XÉT BIỂU ĐỒ THUẬT TOÁN 2
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, LIGHT_BG)
    add_header(s7, "NHẬN XÉT CHI TIẾT: THUẬT TOÁN HISTOGRAM", "Phân tích hiện tượng quá phân đoạn (Over-segmentation) và sai số biên")
    
    metrics2 = [
        ("MAE TRÊN PHONE_F2", "55.0 ms", "RMSE = 57.0 ms", RED),
        ("MAE TOÀN TẬP TEST", "50.0 ms", "Sai số cao nhất trong 3 phương pháp", ORANGE),
        ("SỐ LƯỢNG BIÊN", "2 Đúng / 4 Thừa / 0 Thiếu", "Xuất hiện 4 biên thừa giữa câu", RED)
    ]
    for i, (mtitle, mval, msub, clr) in enumerate(metrics2):
        c = add_card(s7, 0.8 + i * 4.0, 1.6, 3.7, 1.3, WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = mtitle
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        p2 = tf.add_paragraph()
        p2.text = mval
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = clr
        p3 = tf.add_paragraph()
        p3.text = msub
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED

    c_left = add_card(s7, 0.8, 3.2, 5.7, 3.8, WHITE)
    tf = c_left.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "VỊ TRÍ VÀ CHI TIẾT CÁC BIÊN THỪA"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    items_left2 = [
        ("• Biên bắt đầu & kết thúc:", "Bắt đầu 1.09 s (lệch +70 ms); Kết thúc 4.00 s (lệch -40 ms)."),
        ("• Biên thừa tại 2.50 s & 2.80 s:", "STE rơi dưới ngưỡng 0.0572 đủ 200 ms, chia tiếng nói thành Silence giả."),
        ("• Biên thừa tại 3.44 s & 3.64 s:", "Tiếp tục phát sinh thêm 1 đoạn Silence giả nữa trong câu thoại.")
    ]
    for k, v in items_left2:
        p1 = tf.add_paragraph()
        p1.text = k
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = TEAL
        p2 = tf.add_paragraph()
        p2.text = v
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    c_right = add_card(s7, 6.8, 3.2, 5.7, 3.8, WHITE)
    tf = c_right.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "NGUYÊN NHÂN & BIỆN PHÁP CẢI TIẾN"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    items_right2 = [
        ("• Nguyên nhân bản chất:", "Đỉnh Speech trong file phone_F2 nằm ở năng lượng cao, kéo trọng số ngưỡng T lên mức 0.0572, làm các âm tiết yếu bị cắt rời."),
        ("• Ưu điểm bù lại:", "Thuật toán không phụ thuộc dữ liệu huấn luyện, xử lý độc lập và có tính thích nghi khi âm lượng đầu vào thay đổi."),
        ("• Hướng khắc phục:", "Cần tăng trọng số ưu tiên đỉnh Silence (tăng weight) để ép ngưỡng xuống thấp hơn.")
    ]
    for k, v in items_right2:
        p1 = tf.add_paragraph()
        p1.text = k
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = RED
        p2 = tf.add_paragraph()
        p2.text = v
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 8: BIỂU ĐỒ KHỐI THUẬT TOÁN 3 (GAUSSIAN)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, LIGHT_BG)
    add_header(s8, "THUẬT TOÁN 3: THỐNG KÊ GAUSSIAN (STATISTICAL BAYES)", "Sơ đồ khối và mô hình xác suất tối ưu hóa sai số phân lớp")
    
    steps3 = [
        ("BƯỚC 1", "Dữ liệu huấn luyện & LAB", "Tập hợp toàn bộ khung âm thanh từ 4 tệp huấn luyện"),
        ("BƯỚC 2", "Ước lượng (Mean, Std)", "Sil: mean=0.00041, std=0.00092 | Sp: mean=0.20963, std=0.24070"),
        ("BƯỚC 3", "Mô hình hóa mật độ Gauss", "Xây dựng 2 hàm mật độ xác suất p(x|Sil) và p(x|Sp)"),
        ("BƯỚC 4", "Thiết lập phương trình Bayes", "Cân bằng hàm mật độ xác suất dẫn đến phương trình bậc 2: a*x^2 + b*x + c = 0"),
        ("BƯỚC 5", "Giải nghiệm tối ưu", "Tìm nghiệm trong [0, 1] làm cực tiểu tổng lỗi phân lớp tích lũy P(Error)"),
        ("BƯỚC 6", "Khóa ngưỡng T = 0.00358", "Khóa ngưỡng áp dụng cho kiểm thử; loại khoảng lặng ảo < 200 ms")
    ]
    for i, (tag, title, desc) in enumerate(steps3):
        x = 0.8 + (i % 3) * 4.0
        y = 1.7 + (i // 3) * 2.3
        card = add_card(s8, x, y, 3.7, 2.0, WHITE)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ORANGE
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = NAVY
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    banner3 = add_card(s8, 0.8, 6.4, 11.7, 0.65, RGBColor(255, 247, 237), ORANGE)
    tf = banner3.text_frame
    p = tf.paragraphs[0]
    p.text = "CƠ SỞ TOÁN HỌC VỮNG CHẮC: Dựa trên lý thuyết quyết định Bayes, giải nghiệm giải tích chính xác, không dùng thuật toán xấp xỉ."
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY

    # ==========================================
    # SLIDE 9: BIỂU ĐỒ KẾT QUẢ THUẬT TOÁN 3
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9, LIGHT_BG)
    add_header(s9, "KẾT QUẢ THỰC NGHIỆM: GAUSSIAN TRÊN PHONE_F2", "Độ chính xác hoàn hảo — Hai biên dự đoán trùng khít với biên chuẩn")
    
    img_path = Path("ket_qua/phone_F2_statistical.png")
    if img_path.exists():
        s9.shapes.add_picture(str(img_path), Inches(0.8), Inches(1.6), width=Inches(8.5))
        
    side_card3 = add_card(s9, 9.6, 1.6, 2.9, 5.3, WHITE)
    tf = side_card3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "KẾT QUẢ NỔI BẬT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    sections3 = [
        ("Ngưỡng xác suất Bayes", "T = 0.00358 (tách biệt hoàn toàn 2 phân bố)."),
        ("Trùng khít 100%", "Đường xanh liền và đường đỏ đứt trùng nhau hoàn toàn!"),
        ("Không có biên thừa", "Nhiễu nền điện thoại bị triệt tiêu triệt để."),
        ("Không có biên thiếu", "Bắt trọn toàn bộ câu nói từ nguyên âm đầu đến đuôi."),
        ("F0 đồng bộ chuẩn", "Khớp hoàn hảo trong phạm vi 1.02 s đến 4.04 s.")
    ]
    for stitle, sdesc in sections3:
        p1 = tf.add_paragraph()
        p1.text = stitle
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = GREEN
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 10: NHẬN XÉT BIỂU ĐỒ THUẬT TOÁN 3
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10, LIGHT_BG)
    add_header(s10, "NHẬN XÉT CHI TIẾT: THUẬT TOÁN THỐNG KÊ GAUSSIAN", "Phương pháp đạt hiệu quả phân đoạn cao nhất toàn diện")
    
    metrics3 = [
        ("MAE TRÊN PHONE_F2", "0.0 ms", "Sai số tuyệt đối bằng 0", GREEN),
        ("MAE TOÀN TẬP TEST", "12.5 ms", "Tối ưu nhất trong cả 3 phương pháp", GREEN),
        ("SỐ LƯỢNG BIÊN", "2 Đúng / 0 Thừa / 0 Thiếu", "Không phát sinh bất kỳ lỗi biên nào", GREEN)
    ]
    for i, (mtitle, mval, msub, clr) in enumerate(metrics3):
        c = add_card(s10, 0.8 + i * 4.0, 1.6, 3.7, 1.3, WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = mtitle
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        p2 = tf.add_paragraph()
        p2.text = mval
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = clr
        p3 = tf.add_paragraph()
        p3.text = msub
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED

    c_left = add_card(s10, 0.8, 3.2, 5.7, 3.8, WHITE)
    tf = c_left.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "CHI TIẾT SAI LỆCH VỊ TRÍ BIÊN"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    items_left3 = [
        ("• Bắt đầu Speech (1.02 s):", "Chuẩn 1.02 s, dự đoán đúng 1.02 s (lệch đúng 0 ms)."),
        ("• Kết thúc Speech (4.04 s):", "Chuẩn 4.04 s, dự đoán đúng 4.04 s (lệch đúng 0 ms)."),
        ("• Không phát sinh biên thừa:", "Nhiễu ở cuối bản ghi tại 4.76 s có năng lượng nằm dưới 0.00358, bị chặn đứng hoàn toàn.")
    ]
    for k, v in items_left3:
        p1 = tf.add_paragraph()
        p1.text = k
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = GREEN
        p2 = tf.add_paragraph()
        p2.text = v
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    c_right = add_card(s10, 6.8, 3.2, 5.7, 3.8, WHITE)
    tf = c_right.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "LÝ DO ĐẠT KẾT QUẢ VƯỢT TRỘI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    items_right3 = [
        ("• Tận dụng thông tin phân bố:", "Thay vì chỉ dựa vào cận cực trị, mô hình thống kê học được độ phân tán (std) thực tế của năng lượng tiếng nói và khoảng lặng."),
        ("• Vị trí ngưỡng lý tưởng:", "Ngưỡng 0.00358 nằm cao hơn 3.4 lần độ lệch chuẩn Silence (loại sạch nhiễu), nhưng thấp hơn rất nhiều so với trung bình Speech (0.2096)."),
        ("• Khả năng tổng quát hóa:", "Áp dụng ổn định trên mọi môi trường thu âm thử nghiệm (cả Studio và Phone).")
    ]
    for k, v in items_right3:
        p1 = tf.add_paragraph()
        p1.text = k
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = ORANGE
        p2 = tf.add_paragraph()
        p2.text = v
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 11: SO SÁNH TỔNG HỢP 3 THUẬT TOÁN
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11, LIGHT_BG)
    add_header(s11, "SO SÁNH TỔNG HỢP 3 THUẬT TOÁN & ĐÁNH GIÁ KHÁNG NHIỄU", "Bảng đối chiếu toàn diện trên 4 tệp kiểm thử và khảo sát suy giảm SNR")
    
    # Table on Left
    table_shape = s11.shapes.add_table(6, 4, Inches(0.8), Inches(1.6), Inches(6.5), Inches(3.2))
    table = table_shape.table
    table.columns[0].width = Inches(1.7)
    table.columns[1].width = Inches(1.6)
    table.columns[2].width = Inches(1.6)
    table.columns[3].width = Inches(1.6)
    
    table_data = [
        ["Tín hiệu test", "Binary Search", "Histogram", "Gaussian (TK)"],
        ["phone_F2", "45.0 ms (1 thừa)", "55.0 ms (4 thừa)", "0.0 ms (0 lỗi)"],
        ["phone_M2", "10.0 ms", "15.0 ms", "15.0 ms"],
        ["studio_F2", "20.0 ms", "35.0 ms", "25.0 ms"],
        ["studio_M2", "5.0 ms", "95.0 ms", "10.0 ms"],
        ["TRUNG BÌNH GỘP", "20.0 ms", "50.0 ms", "12.5 ms (Tốt nhất)"]
    ]
    for row_idx, row in enumerate(table_data):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            if row_idx == 0:
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            elif row_idx == 5:
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = NAVY
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(220, 252, 231)
            else:
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_MAIN
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE

    # Comparison Image on Right
    comp_img = Path("ket_qua/phone_F2_so_sanh.png")
    if comp_img.exists():
        s11.shapes.add_picture(str(comp_img), Inches(7.6), Inches(1.6), width=Inches(4.9))

    # Bottom summary card
    summary_card = add_card(s11, 0.8, 5.1, 11.7, 1.9, WHITE)
    tf = summary_card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "KẾT LUẬN THỰC NGHIỆM & KHẢO SÁT NHIỄU TRẮNG"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    conclusions = [
        ("• Xếp hạng hiệu năng:", "1. Gaussian (MAE 12.5 ms, 0 lỗi)  >  2. Binary Search (MAE 20.0 ms, 1 lỗi)  >  3. Histogram (MAE 50.0 ms, 4 lỗi)."),
        ("• Khảo sát cộng nhiễu (30 dB -> 0 dB):", "Thuật toán Histogram thích nghi tốt hơn khi SNR suy giảm cực mạnh; trong khi Gaussian & Binary nhạy cảm hơn do dùng ngưỡng cố định.")
    ]
    for k, v in conclusions:
        p1 = tf.add_paragraph()
        p1.text = k + " "
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = CYAN
        # append text
        run = p1.add_run()
        run.text = v
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 12: CHÀO TẠM BIỆT & SẴN SÀNG DEMO
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12, NAVY)
    
    t_box = s12.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.5))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    p.text = "EM XIN CHÂN THÀNH CẢM ƠN THẦY!"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Đề tài đã hoàn thành xuất sắc 100% các yêu cầu kỹ thuật và báo cáo thực nghiệm"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.alignment = PP_ALIGN.CENTER
    
    # 3 Action Cards
    actions = [
        ("SẴN SÀNG DEMO TRỰC TIẾP", "python3 main.py --no-noise\n\nChạy đúng 1 lần duy nhất, tự động mở và dàn đều 4 Figure tại 4 góc màn hình.", CYAN),
        ("KIỂM CHỨNG MÃ NGUỒN", "pytest -q\n\n100% thuật toán tự cài đặt thuần Numpy, không dùng thư viện ngoài Scipy/Toolbox.", GREEN),
        ("KHO LƯU TRỮ GITHUB", "https://github.com/VoThanhQuan-Pentest/XLTHS.git\n\nĐầy đủ toàn bộ code, file nhãn, kết quả và slide báo cáo.", ORANGE)
    ]
    for i, (atitle, adesc, clr) in enumerate(actions):
        card = add_card(s12, 0.8 + i * 4.0, 3.4, 3.7, 2.6, DARK_CARD, clr)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
        p = tf.paragraphs[0]
        p.text = atitle
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = clr
        p2 = tf.add_paragraph()
        p2.text = adesc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        
    f_box = s12.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6))
    tf = f_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Sẵn sàng lắng nghe nhận xét và trả lời câu hỏi phản biện của Giảng viên."
    p.font.size = Pt(13)
    p.font.color.rgb = CYAN
    p.alignment = PP_ALIGN.CENTER

    out_path = Path("Bao_Cao_XLTHS_9_Phut.pptx")
    prs.save(str(out_path))
    print(f"Successfully generated PowerPoint presentation at: {out_path.resolve()}")

if __name__ == "__main__":
    create_deck()
